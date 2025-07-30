"""Microsoft Office document loaders - PowerPoint, Word, Excel."""

from ...core import Attachment, loader
from ... import matchers


@loader(match=matchers.pptx_match)
def pptx_to_python_pptx(att: Attachment) -> Attachment:
    """Load PowerPoint using python-pptx."""
    try:
        from pptx import Presentation
        att._obj = Presentation(att.path)
    except ImportError:
        raise ImportError("python-pptx is required for PowerPoint loading. Install with: pip install python-pptx")
    return att


@loader(match=matchers.docx_match)
def docx_to_python_docx(att: Attachment) -> Attachment:
    """Load Word document using python-docx."""
    try:
        from docx import Document
        att._obj = Document(att.path)
    except ImportError:
        raise ImportError("python-docx is required for Word document loading. Install with: pip install python-docx")
    return att


@loader(match=matchers.excel_match)
def excel_to_openpyxl(att: Attachment) -> Attachment:
    """Load Excel workbook using openpyxl."""
    try:
        from openpyxl import load_workbook
        att._obj = load_workbook(att.path, read_only=True)
    except ImportError:
        raise ImportError("openpyxl is required for Excel loading. Install with: pip install openpyxl")
    return att 